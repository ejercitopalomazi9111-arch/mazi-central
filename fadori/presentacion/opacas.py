"""Ninguna imagen con transparencia: el visor del iPhone las cruza entre láminas.

   python3 opacas.py entrada.pptx salida.pptx

Medido en las capturas de Carlos del 1 de octubre: las fotos JPEG salen SIEMPRE bien; las que se
cruzan (el logo de una lámina en otra, el óvalo de «Pálido» en el lugar del logo) son siempre PNG
con partes transparentes. Así que a cada imagen con transparencia se le pinta por detrás lo que
tiene abajo —el fondo de la lámina y las imágenes de abajo, en su orden— y queda como JPEG opaco
en el mismo sitio. Se ve igual; el visor ya no tiene con qué confundirse.

Si un texto quedaba DEBAJO de una imagen que ahora es opaca y se encima con ella, el texto sube
encima de la imagen (si no, la imagen lo taparía).
"""
import sys, io
import numpy as np
from PIL import Image
from pptx import Presentation
from lxml import etree

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
IN = 914400
DPI = 200                      # resolución del lienzo de la lámina
TOPE = 240                     # px por pulgada máximos de una imagen de salida
q = lambda ns, t: '{%s}%s' % (ns, t)
K = DPI / IN

def blob_de(slide, rid):
    return slide.part.related_part(rid).blob

def fondo_lamina(slide, W, H):
    bg = slide._element.find('.//' + q(P, 'bg'))
    w, h = round(W * K), round(H * K)
    if bg is not None:
        b = bg.find('.//' + q(A, 'blip'))
        if b is not None:
            im = Image.open(io.BytesIO(blob_de(slide, b.get(q(R, 'embed'))))).convert('RGB')
            return im.resize((w, h), Image.LANCZOS)
        c = bg.find('.//' + q(A, 'srgbClr'))
        if c is not None:
            v = c.get('val'); return Image.new('RGB', (w, h), tuple(int(v[i:i + 2], 16) for i in (0, 2, 4)))
    return Image.new('RGB', (w, h), (14, 42, 80))

def caja(el):
    x = el.find('.//' + q(A, 'xfrm'))
    if x is None or x.find(q(A, 'off')) is None: return None
    o, e = x.find(q(A, 'off')), x.find(q(A, 'ext'))
    return int(o.get('x')), int(o.get('y')), int(e.get('cx')), int(e.get('cy'))

def recorte(im, el):
    sr = el.find('.//' + q(A, 'srcRect'))
    if sr is None: return im
    W, H = im.size
    f = lambda k: int(sr.get(k, '0')) / 100000
    l, t, r, b = f('l'), f('t'), f('r'), f('b')
    box = (round(W * l), round(H * t), round(W * (1 - r)), round(H * (1 - b)))
    if box[2] <= box[0] or box[3] <= box[1]: return im
    # recortes negativos (la imagen no llena el marco) se rellenan transparentes
    lienzo = Image.new('RGBA', (box[2] - box[0], box[3] - box[1]), (0, 0, 0, 0))
    lienzo.paste(im.crop((max(0, box[0]), max(0, box[1]), min(W, box[2]), min(H, box[3]))),
                 (max(0, -box[0]), max(0, -box[1])))
    return lienzo

def tiene_alfa(im):
    if im.mode in ('RGBA', 'LA', 'PA') or (im.mode == 'P' and 'transparency' in im.info):
        a = np.asarray(im.convert('RGBA'))[..., 3]
        return a.min() < 250
    return False

def texto_visible(el):
    return el.tag == q(P, 'sp') and ''.join(el.itertext()).strip() != ''

def choca(a, b):
    return a[0] < b[0] + b[2] and b[0] < a[0] + a[2] and a[1] < b[1] + b[3] and b[1] < a[1] + a[3]

def main():
    ent, sal = sys.argv[1:3]
    prs = Presentation(ent); W, H = prs.slide_width, prs.slide_height
    hechas = subidos = 0; vinetas = [0]
    for slide in prs.slides:
        lienzo = fondo_lamina(slide, W, H).convert('RGBA')
        t = slide.shapes._spTree
        por_el = {sh._element: sh for sh in slide.shapes}
        # viñetas de imagen: el visor de Apple tampoco las quiere; viñeta de letra en su lugar
        for bb in list(t.iter(q(A, 'buBlip'))):
            pPr = bb.getparent(); pos = list(pPr).index(bb); pPr.remove(bb)
            for bf in pPr.findall(q(A, 'buFont')): bf.set('typeface', 'Arial'); [bf.attrib.pop(k, None) for k in ('pitchFamily', 'charset')]
            pPr.insert(pos, etree.Element(q(A, 'buChar'), char='•')); vinetas[0] += 1
        i = 0
        while i < len(t):
            el = t[i]; i += 1
            if el.tag != q(P, 'pic'): continue
            b = caja(el)
            if b is None and el in por_el and por_el[el].left is not None:      # marcador: hereda la posición
                sh = por_el[el]; b = (sh.left, sh.top, sh.width, sh.height)
            blip = el.find('.//' + q(A, 'blip'))
            if b is None or blip is None or b[2] <= 0 or b[3] <= 0: continue
            rid = blip.get(q(R, 'embed'))
            try:
                im = Image.open(io.BytesIO(blob_de(slide, rid)))
                im.load(); fmt = im.format; original = im
            except Exception:
                continue
            im = recorte(im, el); im.format = fmt
            # tamaño de salida: el del marco a TOPE ppp como máximo, sin inventar pixeles de más
            fw, fh = b[2] / IN, b[3] / IN
            ow = max(8, min(round(fw * TOPE), max(im.size[0], round(fw * 96))))
            oh = max(8, round(ow * fh / fw))
            px = (round(b[0] * K), round(b[1] * K), round((b[0] + b[2]) * K), round((b[1] + b[3]) * K))
            # lo de abajo, también donde el marco se sale de la lámina
            abajo = Image.new('RGBA', (px[2] - px[0], px[3] - px[1]), (14, 42, 80, 255))
            abajo.paste(lienzo.crop((max(0, px[0]), max(0, px[1]), min(lienzo.width, px[2]), min(lienzo.height, px[3]))),
                        (max(0, -px[0]), max(0, -px[1])))
            if not tiene_alfa(original) and (original.mode != 'RGB' or (fmt or '').upper() not in ('JPEG', 'MPO')):
                # opaca pero no JPEG: se pasa a JPEG ENTERA, con su recorte intacto (portadas de video, etc.)
                buf = io.BytesIO(); original.convert('RGB').save(buf, 'JPEG', quality=92, subsampling=0)
                _, nuevo = slide.part.get_or_add_image_part(io.BytesIO(buf.getvalue()))
                blip.set(q(R, 'embed'), nuevo); hechas += 1
                pinta = im.convert('RGBA').resize((ow, oh), Image.LANCZOS)
            elif tiene_alfa(im) or tiene_alfa(original):
                arriba = im.convert('RGBA').resize((ow, oh), Image.LANCZOS)
                res = abajo.resize((ow, oh), Image.LANCZOS); res.alpha_composite(arriba)
                buf = io.BytesIO(); res.convert('RGB').save(buf, 'JPEG', quality=90, subsampling=0)
                _, nuevo = slide.part.get_or_add_image_part(io.BytesIO(buf.getvalue()))
                blip.set(q(R, 'embed'), nuevo)
                for c in list(blip): blip.remove(c)                       # sin efectos de color encima
                sr = el.find('.//' + q(A, 'srcRect'))
                if sr is not None:
                    for k in list(sr.attrib): del sr.attrib[k]
                hechas += 1
                # el texto que estaba abajo y se encima sube arriba de la imagen
                pos = list(t).index(el)
                for prev in list(t)[:pos]:
                    if texto_visible(prev) or (prev.tag == q(P, 'sp') and prev.find('.//' + q(A, 'solidFill')) is not None and prev.find(q(P, 'spPr')).find(q(A, 'solidFill')) is not None):
                        cb = caja(prev)
                        if cb is None and prev in por_el and por_el[prev].left is not None:
                            sh = por_el[prev]; cb = (sh.left, sh.top, sh.width, sh.height)
                        if cb and choca(cb, b):
                            t.remove(prev); t.insert(list(t).index(el) + 1, prev); subidos += 1
                pinta = res
            else:
                pinta = im.convert('RGBA').resize((ow, oh), Image.LANCZOS)
            # se pinta en el lienzo para lo que venga encima
            pinta = pinta.resize((px[2] - px[0], px[3] - px[1]), Image.LANCZOS).convert('RGBA')
            x0, y0 = max(0, px[0]), max(0, px[1]); x1, y1 = min(lienzo.width, px[2]), min(lienzo.height, px[3])
            if x1 > x0 and y1 > y0:
                trozo = pinta.crop((x0 - px[0], y0 - px[1], x1 - px[0], y1 - px[1]))
                lienzo.alpha_composite(trozo, (x0, y0))
    # PowerPoint exige el orden del esquema dentro del texto (la viñeta antes de
    # defRPr, etc.); LibreOffice no. Sin esto, Brigada y Acuática no abrían.
    from orden import ordenar
    for s_ in prs.slides: ordenar(s_._element)
    prs.save(sal)
    print(f'opacas: {hechas} imágenes con transparencia vueltas opacas · {subidos} textos subidos encima · {vinetas[0]} viñetas de imagen cambiadas')

if __name__ == '__main__': main()
