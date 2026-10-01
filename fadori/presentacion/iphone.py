"""Que la vista previa del iPhone (y la de WhatsApp, que es la misma) pinte la presentación como es.

   python3 iphone.py entrada.pptx salida.pptx

El visor de Apple (OfficeImport) sólo dibuja con HTML el rectángulo simple. Todo lo demás —esquinas
redondeadas, óvalos, flechas, líneas, grupos, sombras, giros— lo vuelve un PDF aparte, y esa memoria
de PDFs tiene un error conocido: se cruza entre láminas (la tarjeta de la portada aparece en la 3, la
foto del pie en tres láminas). Y unas figuras ni las dibuja: round2DiagRect, donut, moon, wave…

Aquí cada una de esas figuras se HORNEA: LibreOffice la pinta sola sobre negro y sobre blanco, de la
diferencia sale la transparencia exacta, y en su lugar queda una imagen rectangular en el mismo sitio.
Si la figura traía texto, el texto se queda como cuadro de texto rectangular encima: se sigue editando.
"""
import sys, io, os, copy, math, subprocess, tempfile, shutil
import numpy as np
from PIL import Image
from pptx import Presentation
from lxml import etree

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
IN = 914400
DPI = 200
q = lambda ns, t: '{%s}%s' % (ns, t)

# ─────────────────────────── qué se hornea ───────────────────────────
def spPr(el):
    for t in ('spPr', 'grpSpPr'):
        e = el.find(q(P, t))
        if e is not None: return e
    return None

def xfrm(el):
    sp = spPr(el)
    return sp.find(q(A, 'xfrm')) if sp is not None else None

def geometria(el):
    sp = spPr(el)
    if sp is None: return None
    g = sp.find(q(A, 'prstGeom'))
    if g is not None: return g.get('prst')
    if sp.find(q(A, 'custGeom')) is not None: return 'custGeom'
    return None

def con_efectos(el):
    sp = spPr(el)
    if sp is None: return False
    e = sp.find(q(A, 'effectLst'))
    return (e is not None and len(e)) or sp.find(q(A, 'effectDag')) is not None

def girada(el):
    x = xfrm(el)
    return x is not None and (int(x.get('rot', '0')) % 21600000 != 0 or x.get('flipH') == '1' or x.get('flipV') == '1')

def degradado_largo(el):
    sp = spPr(el)
    g = sp.find(q(A, 'gradFill')) if sp is not None else None
    return g is not None and len(g.findall('.//' + q(A, 'gs'))) >= 3

def tiene_texto(el):
    return el.tag == q(P, 'sp') and ''.join(el.itertext()).strip() != '' and el.find(q(P, 'txBody')) is not None

def hay_que_hornear(el):
    t = el.tag
    if t in (q(P, 'cxnSp'), q(P, 'grpSp')): return True
    if t not in (q(P, 'sp'), q(P, 'pic')): return False
    g = geometria(el)
    return (g not in (None, 'rect')) or con_efectos(el) or girada(el) or degradado_largo(el)

# ─────────────────────────── desagrupar ───────────────────────────
def desagrupar(spTree):
    """Los grupos sin giro se deshacen: cada hijo baja a la lámina con sus coordenadas reales."""
    hechos = 0
    cambio = True
    while cambio:
        cambio = False
        for g in list(spTree):
            if g.tag != q(P, 'grpSp'): continue
            gx = g.find(q(P, 'grpSpPr')).find(q(A, 'xfrm'))
            if gx is None or girada(g): continue
            off, ext = gx.find(q(A, 'off')), gx.find(q(A, 'ext'))
            choff, chext = gx.find(q(A, 'chOff')), gx.find(q(A, 'chExt'))
            if None in (off, ext, choff, chext): continue
            ox, oy, ew, eh = (int(off.get('x')), int(off.get('y')), int(ext.get('cx')), int(ext.get('cy')))
            cx, cy, cw, ch = (int(choff.get('x')), int(choff.get('y')), int(chext.get('cx')), int(chext.get('cy')))
            sx = ew / cw if cw else 1; sy = eh / ch if ch else 1
            pos = list(spTree).index(g)
            hijos = [h for h in g if h.tag not in (q(P, 'nvGrpSpPr'), q(P, 'grpSpPr'))]
            for h in hijos:
                hx = xfrm(h)
                if hx is not None:
                    o, e = hx.find(q(A, 'off')), hx.find(q(A, 'ext'))
                    if o is not None:
                        o.set('x', str(round(ox + (int(o.get('x')) - cx) * sx))); o.set('y', str(round(oy + (int(o.get('y')) - cy) * sy)))
                    if e is not None:
                        e.set('cx', str(round(int(e.get('cx')) * sx))); e.set('cy', str(round(int(e.get('cy')) * sy)))
                    if h.tag == q(P, 'grpSp'):
                        co, ce = hx.find(q(A, 'chOff')), hx.find(q(A, 'chExt'))   # el hijo-grupo conserva su espacio interno
                    # letra: si el grupo escalaba, el texto también
                if abs(sx - 1) > .02 or abs(sy - 1) > .02:
                    s = min(sx, sy)
                    for r in h.iter(q(A, 'rPr'), q(A, 'defRPr'), q(A, 'endParaRPr')):
                        if r.get('sz'): r.set('sz', str(max(600, round(int(r.get('sz')) * s))))
                g.remove(h); spTree.insert(pos, h); pos += 1
            spTree.remove(g); hechos += 1; cambio = True
    return hechos

# ─────────────────────────── dónde está ───────────────────────────
def caja(el, shape_por_id):
    x = xfrm(el)
    if x is not None and x.find(q(A, 'off')) is not None:
        o, e = x.find(q(A, 'off')), x.find(q(A, 'ext'))
        return int(o.get('x')), int(o.get('y')), int(e.get('cx')), int(e.get('cy'))
    sh = shape_por_id.get(id_de(el))       # placeholder: hereda la posición del diseño
    if sh is not None and sh.left is not None: return sh.left, sh.top, sh.width, sh.height
    return None

def id_de(el):
    c = el.find('.//' + q(P, 'cNvPr'))
    return c.get('id') if c is not None else None

def fijar_caja(el, b):
    sp = spPr(el)
    x = sp.find(q(A, 'xfrm'))
    if x is None:
        x = etree.Element(q(A, 'xfrm')); sp.insert(0, x)
    for c in list(x): x.remove(c)
    etree.SubElement(x, q(A, 'off'), x=str(b[0]), y=str(b[1])); etree.SubElement(x, q(A, 'ext'), cx=str(b[2]), cy=str(b[3]))

# ─────────────────────────── pintar con LibreOffice ───────────────────────────
def sin_texto(el):
    tb = el.find(q(P, 'txBody'))
    if tb is None: return
    for p in tb.findall(q(A, 'p')): tb.remove(p)
    etree.SubElement(tb, q(A, 'p'))

def sin_placeholder(el, b):
    for ph in el.iter(q(P, 'ph')): ph.getparent().remove(ph)
    if b is not None and xfrm(el) is None: fijar_caja(el, b)

def fondo(slide, color):
    cSld = slide._element.find(q(P, 'cSld'))
    bg = cSld.find(q(P, 'bg'))
    if bg is not None: cSld.remove(bg)
    bg = etree.SubElement(cSld, q(P, 'bg')); cSld.remove(bg); cSld.insert(0, bg)
    bp = etree.SubElement(bg, q(P, 'bgPr'))
    sf = etree.SubElement(bp, q(A, 'solidFill')); etree.SubElement(sf, q(A, 'srgbClr'), val=color)
    etree.SubElement(bp, q(A, 'effectLst'))
    slide._element.set('showMasterSp', '0')

def copiar_rels(el, de, a):
    """Las imágenes y rellenos que la figura cita, a la lámina de prueba."""
    for n in el.iter():
        for k, v in list(n.attrib.items()):
            if not k.startswith('{%s}' % R): continue
            if n.tag in (q(A, 'hlinkClick'), q(A, 'hlinkHover')):
                n.getparent().remove(n); break
            try:
                rel = de.part.rels[v]
            except KeyError:
                continue
            if rel.is_external:
                n.attrib[k] = a.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            else:
                n.attrib[k] = a.part.relate_to(rel.target_part, rel.reltype)

def pintar(ent, trabajos, carpeta):
    """trabajos: [(índice de lámina, copia del elemento)] → dos PDF, uno sobre negro y otro sobre blanco."""
    paginas = {}
    for color in ('000000', 'FFFFFF'):
        prs = Presentation(ent)
        orig = list(prs.slides)
        for i, el in trabajos:
            s = orig[i]
            nueva = prs.slides.add_slide(s.slide_layout)
            for ph in list(nueva.placeholders): ph._element.getparent().remove(ph._element)
            fondo(nueva, color)
            c = copy.deepcopy(el)
            copiar_rels(c, s, nueva)
            nueva.shapes._spTree.append(c)
        lst = prs.slides._sldIdLst
        for sid in list(lst)[:len(orig)]: lst.remove(sid)
        f = os.path.join(carpeta, f'p{color}.pptx'); prs.save(f)
        subprocess.run(['soffice', '--headless', '--norestore', f'-env:UserInstallation=file://{carpeta}/ui',
                        '--convert-to', 'pdf', '--outdir', carpeta, f], check=True, capture_output=True, timeout=900)
        pdf = f[:-5] + '.pdf'
        subprocess.run(['pdftoppm', '-png', '-r', str(DPI), pdf, os.path.join(carpeta, f'r{color}')], check=True)
        hojas = sorted(n for n in os.listdir(carpeta) if n.startswith(f'r{color}-'))
        if len(hojas) != len(trabajos): raise SystemExit(f'iphone: esperaba {len(trabajos)} páginas y salieron {len(hojas)}')
        paginas[color] = [os.path.join(carpeta, n) for n in hojas]
    return paginas

def recortar(negro, blanco, b, W, H):
    """Transparencia exacta: lo que cambia entre fondo negro y blanco es lo que deja pasar."""
    m = int(.25 * IN)
    x0, y0 = max(0, b[0] - m), max(0, b[1] - m)
    x1, y1 = min(W, b[0] + b[2] + m), min(H, b[1] + b[3] + m)
    k = DPI / IN
    px = lambda v: int(round(v * k))
    bx = (px(x0), px(y0), px(x1), px(y1))
    N = np.asarray(Image.open(negro).convert('RGB').crop(bx), dtype=np.float32)
    B = np.asarray(Image.open(blanco).convert('RGB').crop(bx), dtype=np.float32)
    alfa = 1 - (B - N).mean(axis=2) / 255
    alfa = np.clip(alfa, 0, 1)
    color = np.where(alfa[..., None] > 1e-3, N / np.maximum(alfa[..., None], 1e-3), 0)
    rgba = np.dstack([np.clip(color, 0, 255), alfa * 255]).round().astype(np.uint8)
    rgba[alfa < 3 / 255] = 0
    ys, xs = np.nonzero(rgba[..., 3] > 2)
    if not len(xs): return None, None
    ya, yb, xa, xb = ys.min(), ys.max() + 1, xs.min(), xs.max() + 1
    rgba = rgba[ya:yb, xa:xb]
    rgba[rgba[..., 3] >= 252, 3] = 255
    caja_emu = (x0 + round(xa / k), y0 + round(ya / k), round((xb - xa) / k), round((yb - ya) / k))
    return Image.fromarray(rgba, 'RGBA'), caja_emu

def lamina_completa(ent, carpeta):
    """La presentación tal cual, para saber qué hay DEBAJO de cada esquina."""
    d = os.path.join(carpeta, 'completa'); os.makedirs(d, exist_ok=True)
    subprocess.run(['soffice', '--headless', '--norestore', f'-env:UserInstallation=file://{carpeta}/ui',
                    '--convert-to', 'pdf', '--outdir', d, ent], check=True, capture_output=True, timeout=900)
    pdf = os.path.join(d, os.path.basename(ent)[:-5] + '.pdf')
    subprocess.run(['pdftoppm', '-png', '-r', str(DPI), pdf, os.path.join(d, 'l')], check=True)
    return [os.path.join(d, n) for n in sorted(os.listdir(d)) if n.startswith('l-')]

def esquinas(negro, blanco, completa, b):
    """Foto con esquinas redondeadas: la foto se queda como es (JPEG, rectangular) y encima va SÓLO lo
    que cubren sus esquinas, tomado de la lámina completa (la tarjeta blanca o el fondo)."""
    k = DPI / IN
    bx = tuple(int(round(v * k)) for v in (b[0], b[1], b[0] + b[2], b[1] + b[3]))
    N = np.asarray(Image.open(negro).convert('RGB').crop(bx), dtype=np.float32)
    B = np.asarray(Image.open(blanco).convert('RGB').crop(bx), dtype=np.float32)
    alfa = np.clip(1 - (B - N).mean(axis=2) / 255, 0, 1)
    tapa = 1 - alfa
    tapa[tapa < 4 / 255] = 0
    if tapa.max() == 0: return None
    C = np.asarray(Image.open(completa).convert('RGB').crop(bx), dtype=np.uint8)
    rgba = np.dstack([C, (tapa * 255).round().astype(np.uint8)])
    rgba[rgba[..., 3] == 0] = 0          # lo transparente sin color: si no, el PNG carga la foto entera
    return Image.fromarray(rgba, 'RGBA')

def a_bytes(im, foto_original=None):
    a = np.asarray(im)[..., 3]
    buf = io.BytesIO()
    if a.min() >= 250:
        im.convert('RGB').save(buf, 'JPEG', quality=88); return buf.getvalue(), '.jpg'
    im.save(buf, 'PNG', optimize=True); return buf.getvalue(), '.png'

# ─────────────────────────── la imagen y el texto que quedan ───────────────────────────
def pic_xml(id_, nombre, rid, b):
    x = f'''<p:pic xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}"><p:nvPicPr><p:cNvPr id="{id_}" name="{nombre}"/><p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr/></p:nvPicPr><p:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill><p:spPr><a:xfrm><a:off x="{b[0]}" y="{b[1]}"/><a:ext cx="{b[2]}" cy="{b[3]}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'''
    return etree.fromstring(x)

def margen_inscrito(geo, el, b):
    """El texto de una figura no rectangular vive en su rectángulo inscrito; al volverla rectángulo se respeta."""
    w, h = b[2], b[3]
    adj = None
    g = spPr(el).find(q(A, 'prstGeom')) if spPr(el) is not None else None
    if g is not None:
        for gd in g.iter(q(A, 'gd')):
            f = gd.get('fmla', '')
            if f.startswith('val '): adj = int(f[4:]); break
    if geo == 'roundRect':
        r = (adj if adj is not None else 16667) / 100000 * min(w, h)
        d = round(r * .29289); return d, d, d, d
    if geo == 'round2DiagRect':
        r = (adj if adj is not None else 16667) / 100000 * min(w, h)
        d = round(r * .29289); return d, d, d, d
    if geo == 'ellipse':
        dx, dy = round(w * .14645), round(h * .14645); return dx, dy, dx, dy
    return 0, 0, 0, 0

def texto_rectangular(el, b, geo):
    sp = spPr(el)
    for t in ('prstGeom', 'custGeom', 'noFill', 'solidFill', 'gradFill', 'blipFill', 'pattFill', 'grpFill', 'ln', 'effectLst', 'effectDag', 'scene3d', 'sp3d'):
        for n in sp.findall(q(A, t)): sp.remove(n)
    fijar_caja(el, b)
    x = sp.find(q(A, 'xfrm'))
    pos = list(sp).index(x) + 1
    g = etree.Element(q(A, 'prstGeom'), prst='rect'); etree.SubElement(g, q(A, 'avLst'))
    sp.insert(pos, g)
    sp.insert(pos + 1, etree.Element(q(A, 'noFill')))
    ln = etree.Element(q(A, 'ln')); etree.SubElement(ln, q(A, 'noFill')); sp.insert(pos + 2, ln)
    sp.insert(pos + 3, etree.Element(q(A, 'effectLst')))
    bp = el.find(q(P, 'txBody')).find(q(A, 'bodyPr'))
    dl, dt, dr, db = margen_inscrito(geo, el, b)
    for k, d, base in (('lIns', dl, 91440), ('tIns', dt, 45720), ('rIns', dr, 91440), ('bIns', db, 45720)):
        if d: bp.set(k, str(int(bp.get(k, base)) + d))

def nuevo_id(spTree):
    ids = [int(c.get('id')) for c in spTree.iter(q(P, 'cNvPr')) if c.get('id', '').isdigit()]
    return max(ids + [1]) + 1

def main():
    ent, sal = sys.argv[1:3]
    tmp = tempfile.mkdtemp(prefix='iphone-')
    # 1 · desagrupar y arreglar ids raros (id 0 o repetidos dentro de la lámina)
    prs = Presentation(ent)
    deshechos = 0
    for s in prs.slides:
        t = s.shapes._spTree
        deshechos += desagrupar(t)
        vistos = set()
        for c in t.iter(q(P, 'cNvPr')):
            if c.getparent().getparent() is t and False: pass
            i = c.get('id')
            if i in (None, '0') or i in vistos:
                c.set('id', str(nuevo_id(t)))
            vistos.add(c.get('id'))
    paso1 = os.path.join(tmp, 'paso1.pptx'); prs.save(paso1)
    # 2 · qué se hornea
    prs = Presentation(paso1)
    W, H = prs.slide_width, prs.slide_height
    trabajos, meta = [], []
    for i, s in enumerate(prs.slides):
        por_id = {str(sh.shape_id): sh for sh in s.shapes}
        for el in list(s.shapes._spTree):
            if not hay_que_hornear(el): continue
            b = caja(el, por_id)
            if b is None or b[2] <= 0 and b[3] <= 0: continue
            geo = geometria(el)
            texto = tiene_texto(el) and not girada(el)
            modo = 'esquinas' if (el.tag == q(P, 'pic') and not con_efectos(el) and not girada(el)) else 'hornear'
            c = copy.deepcopy(el)
            sin_placeholder(c, b)
            if texto: sin_texto(c)
            trabajos.append((i, c)); meta.append((i, id_de(el), b, geo, texto, modo))
    if not trabajos:
        shutil.copy(paso1, sal); print('iphone: nada que hornear'); return
    paginas = pintar(paso1, trabajos, tmp)
    completas = lamina_completa(paso1, tmp) if any(m[5] == 'esquinas' for m in meta) else []
    # 3 · cambiar cada figura por su imagen (y su texto encima)
    slides = list(prs.slides)
    hechas = quitadas = 0; peso = 0
    for k, (i, idv, b, geo, texto, modo) in enumerate(meta):
        s = slides[i]; t = s.shapes._spTree
        el = next((e for e in t if id_de(e) == idv), None)
        if el is None: continue
        if modo == 'esquinas':
            g = spPr(el).find(q(A, 'prstGeom'))
            if g is not None:
                g.set('prst', 'rect')
                for c in list(g): g.remove(c)
                etree.SubElement(g, q(A, 'avLst'))
            else:
                for c in spPr(el).findall(q(A, 'custGeom')): spPr(el).remove(c)
            im = esquinas(paginas['000000'][k], paginas['FFFFFF'][k], completas[i], b)
            if im is not None:
                datos, _ = a_bytes(im); peso += len(datos)
                _, rid = s.part.get_or_add_image_part(io.BytesIO(datos))
                t.insert(list(t).index(el) + 1, pic_xml(nuevo_id(t), 'Esquinas (horneadas)', rid, b)); hechas += 1
            continue
        im, bx = recortar(paginas['000000'][k], paginas['FFFFFF'][k], b, W, H)
        pos = list(t).index(el)
        if im is not None:
            datos, ext = a_bytes(im)
            peso += len(datos)
            _, rid = s.part.get_or_add_image_part(io.BytesIO(datos))
            nombre = (el.find('.//' + q(P, 'cNvPr')).get('name') or 'Figura') + ' (horneada)'
            pid = nuevo_id(t) if texto else int(idv)
            if not texto: t.remove(el)
            t.insert(pos, pic_xml(pid, nombre, rid, bx)); hechas += 1
        elif not texto:
            t.remove(el); quitadas += 1
        if texto:
            texto_rectangular(el, b, geo)
    prs.save(sal)
    shutil.rmtree(tmp, ignore_errors=True)
    print(f'iphone: {deshechos} grupos deshechos · {hechas} figuras horneadas ({peso // 1024} KB) · {quitadas} invisibles quitadas')

if __name__ == '__main__': main()
