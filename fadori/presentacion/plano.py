"""Versión «para ver en el celular»: cada lámina es UNA imagen de fondo y nada más.

   python3 plano.py entrada.pptx salida.pptx salida.pdf

La vista previa de Apple (iPhone, WhatsApp, Mail) guarda lo que ya pintó con la dirección de memoria
de cada elemento y, al liberar una lámina, la siguiente reusa esas direcciones: un elemento nuevo
recibe la imagen de uno viejo. Sin elementos en la lámina no hay a quién darle una imagen ajena.
Además sale el PDF, que el iPhone y WhatsApp abren siempre bien.
"""
import sys, io, os, subprocess, tempfile, shutil
from pptx import Presentation
from pptx.util import Emu
from lxml import etree

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
q = lambda ns, t: '{%s}%s' % (ns, t)

def main():
    ent, sal, pdf_sal = sys.argv[1:4]
    tmp = tempfile.mkdtemp(prefix='plano-')
    subprocess.run(['soffice', '--headless', '--norestore', f'-env:UserInstallation=file://{tmp}/ui',
                    '--convert-to', 'pdf', '--outdir', tmp, ent], check=True, capture_output=True, timeout=1200)
    pdf = os.path.join(tmp, os.path.basename(ent)[:-5] + '.pdf')
    shutil.copy(pdf, pdf_sal)
    orig = Presentation(ent)
    W, H = orig.slide_width, orig.slide_height
    ancho = 1920
    dpi = round(ancho / (W / 914400))
    subprocess.run(['pdftoppm', '-jpeg', '-jpegopt', 'quality=88', '-r', str(dpi), pdf, os.path.join(tmp, 'l')], check=True)
    hojas = sorted(n for n in os.listdir(tmp) if n.startswith('l-') and n.endswith('.jpg'))
    if len(hojas) != len(orig.slides): raise SystemExit(f'plano: {len(hojas)} páginas para {len(orig.slides)} láminas')
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    vacio = prs.slide_layouts[6]
    for n, (h, so) in enumerate(zip(hojas, orig.slides)):
        s = prs.slides.add_slide(vacio)
        for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
        with open(os.path.join(tmp, h), 'rb') as f:
            _, rid = s.part.get_or_add_image_part(io.BytesIO(f.read()))
        cSld = s._element.find(q(P, 'cSld'))
        bg = etree.Element(q(P, 'bg')); bp = etree.SubElement(bg, q(P, 'bgPr'))
        bf = etree.SubElement(bp, q(A, 'blipFill'), dpi='0', rotWithShape='1')
        etree.SubElement(bf, q(A, 'blip')).set(q(R, 'embed'), rid)
        etree.SubElement(bf, q(A, 'srcRect')); st = etree.SubElement(bf, q(A, 'stretch')); etree.SubElement(st, q(A, 'fillRect'))
        etree.SubElement(bp, q(A, 'effectLst'))
        cSld.insert(0, bg)
        s._element.set('showMasterSp', '0')
        # las notas del orador se conservan
        if so.has_notes_slide and so.notes_slide.notes_text_frame is not None:
            t = so.notes_slide.notes_text_frame.text
            if t.strip(): s.notes_slide.notes_text_frame.text = t
    prs.save(sal)
    shutil.rmtree(tmp, ignore_errors=True)
    print(f'plano: {len(hojas)} láminas · {os.path.getsize(sal) // 1024} KB · PDF {os.path.getsize(pdf_sal) // 1024} KB')

if __name__ == '__main__': main()
