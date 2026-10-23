"""formato.py · «Mis patrones saludables» igual a la hoja que mandaron (Lección 2, Quantrix).

Carlos, 22 de octubre: «hazme una hoja de Excel idéntica a este formato, ya no
hagas el forms». Dos pestañas:
  · «Patrones v1» — el formato, cuadro por cuadro: el marco azul, el título, la
                    tabla de promedio/máximo/mínimo, la comparativa por nivel de
                    energía, la gráfica de barras y el patrón. Todo con fórmulas.
  · «Datos»       — donde se anotan las respuestas del grupo, una fila por persona.

Trae las 10 respuestas del Forms que mandó Carlos en captura (VACIO=1 la saca
sin ellas). Sin datos, las celdas quedan vacías en vez de enseñar ceros.

  python3.13 escuela/patrones/formato.py  →  escuela/patrones/mis-patrones-saludables.xlsx
"""
import os
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
SALE = os.path.join(AQUI, 'mis-patrones-saludables.xlsx')
N = 200                                         # cupo de respuestas en «Datos»

AZUL, AZUL_TITULO, NEGRO, GRIS = '3D6FE0', '1E78D7', '1F2328', '6B7280'
FUENTE = 'Arial'
linea = Side(style='thin', color=NEGRO)
gruesa = Side(style='medium', color=NEGRO)
caja = Border(left=linea, right=linea, top=linea, bottom=linea)
centro = Alignment(horizontal='center', vertical='center', wrap_text=True)
izquierda = Alignment(horizontal='left', vertical='center', wrap_text=True, indent=1)

def F(tam=11, negrita=False, color=NEGRO, italica=False):
    return Font(name=FUENTE, size=tam, bold=negrita, color=color, italic=italica)

def marco(ws, rango, borde=gruesa):
    """Contorno de un rango (sin tocar las líneas de adentro)."""
    filas = list(ws[rango])
    for i, fila in enumerate(filas):
        for j, c in enumerate(fila):
            b = c.border
            c.border = Border(left=borde if j == 0 else b.left, right=borde if j == len(fila) - 1 else b.right,
                              top=borde if i == 0 else b.top, bottom=borde if i == len(filas) - 1 else b.bottom)

def rejilla(ws, rango):
    for fila in ws[rango]:
        for c in fila: c.border = caja

wb = Workbook()

# ══ Datos ═════════════════════════════════════════════════════════════════
# Igual que la hoja de respuestas del Forms (captura de Carlos, 22 de octubre):
# marca temporal, vasos de agua, horas de sueño y horas de actividad. Ahí la gente
# contestó a veces «3 vasos», «9 horas xd» o «1 o 2 horas», así que A–E guardan lo
# que se pegue tal cual y G–I sacan el número que de verdad se usa (con «1 o 2»
# toma el punto medio). El Forms no preguntó el nivel de energía: E se llena a mano.
da = wb.active; da.title = 'Datos'
cab = ['Marca temporal', 'Vasos de agua que tomas', 'Horas de sueño promedio', 'Horas de actividad física',
       'Nivel de energía', None, 'Agua (número)', 'Sueño (número)', 'Actividad (número)']
for i, t in enumerate(cab, start=1):
    if t is None: continue
    c = da.cell(row=1, column=i, value=t)
    gris = i >= 7
    c.font = F(11, True, NEGRO if gris else 'FFFFFF')
    c.fill = PatternFill('solid', fgColor='E5E7EB' if gris else AZUL); c.alignment = centro; c.border = caja
for col, w in zip('ABCDEFGHI', [21, 17, 17, 17, 15, 3, 13, 13, 13]): da.column_dimensions[col].width = w
da.row_dimensions[1].height = 34
da.freeze_panes = 'A2'
da['G' + str(N + 3)] = 'G, H e I salen solas: sacan el número de lo que se escribió («3 vasos» → 3, «1 o 2 horas» → 1.5).'
da['G' + str(N + 3)].font = F(9, color=GRIS, italica=True)

def numero(ref):
    """El número de una respuesta: 5 → 5, «3 vasos» → 3, «1 o 2 horas» → 1.5, vacío → ''."""
    x = f'TRIM({ref})'
    o = f'SEARCH(" o ",{x})'
    antes = f'VALUE(LEFT({x},{o}-1))'
    despues = f'VALUE(LEFT(MID({x},{o}+3,99),FIND(" ",MID({x},{o}+3,99)&" ")-1))'
    primero = f'VALUE(LEFT({x},FIND(" ",{x}&" ")-1))'
    return (f'=IF({x}="","",IFERROR(VALUE({ref}),IFERROR(({antes}+{despues})/2,IFERROR({primero},""))))')

for r in range(2, N + 2):
    for col in 'ABCDE':
        da[f'{col}{r}'].font = F(); da[f'{col}{r}'].alignment = centro
        if r < 42: da[f'{col}{r}'].border = caja
    da[f'A{r}'].number_format = 'dd/mm/yyyy hh:mm:ss'
    for src, dst in zip('BCD', 'GHI'):
        da[f'{dst}{r}'] = numero(f'{src}{r}')
        da[f'{dst}{r}'].font = F(color=GRIS); da[f'{dst}{r}'].alignment = centro
        if r < 42: da[f'{dst}{r}'].border = caja
energia = DataValidation(type='list', formula1='"Bajo,Medio,Alto"', allow_blank=True,
                         errorTitle='Nivel de energía', error='Elige Bajo, Medio o Alto.', showErrorMessage=True)
da.add_data_validation(energia); energia.add(f'E2:E{N + 1}')
da.print_area = 'A1:E41'
da.page_setup.fitToWidth = 1; da.page_setup.fitToHeight = 0
da.sheet_properties.pageSetUpPr.fitToPage = True

# las respuestas que ya llegaron, copiadas de la captura tal como se contestaron
from datetime import datetime
RESPUESTAS = [
    ('12:45:53', 1, 4, 4), ('12:46:32', 5, 6, 8), ('12:46:32', 8, 10, 3),
    ('12:46:50', '3 vasos', '7 horas', '1 hora'), ('12:46:55', 4, 8, 2), ('12:47:14', 7, 3, 7),
    ('12:47:35', 9, 7, 1), ('12:47:50', '5 vasos', '6 horas', '1 hora'), ('12:48:19', 5, 6, 1),
    ('12:48:32', 4, '9 horas xd', '1 o 2 horas'),
]
if os.environ.get('VACIO') != '1':
    for i, (hora, agua, sueno, act) in enumerate(RESPUESTAS):
        r = 2 + i
        da[f'A{r}'] = datetime.strptime('22/10/2026 ' + hora, '%d/%m/%Y %H:%M:%S')
        da[f'B{r}'], da[f'C{r}'], da[f'D{r}'] = agua, sueno, act

def R(col): return f"Datos!${col}$2:${col}${N + 1}"

# ══ Patrones v1 · el formato ══════════════════════════════════════════════
ws = wb.create_sheet('Patrones v1')
wb.active = 1
ws.sheet_view.showGridLines = False
# A y L son el marco azul; B y K el aire entre el marco y los cuadros; C..J el papel
anchos = {'A': 3.5, 'B': 2.5, 'C': 6.5, 'D': 15, 'E': 13, 'F': 7, 'G': 11, 'H': 11, 'I': 11, 'J': 4, 'K': 2.5, 'L': 3.5}
for col, w in anchos.items(): ws.column_dimensions[col].width = w
ULTIMA = 27
azul = PatternFill('solid', fgColor=AZUL)
blanco = PatternFill('solid', fgColor='FFFFFF')
for r in range(1, ULTIMA + 1):
    for col in 'ABCDEFGHIJKL':
        ws[f'{col}{r}'].fill = azul if col in 'AL' or r in (1, ULTIMA) else blanco
for r, h in {1: 12, ULTIMA: 12}.items(): ws.row_dimensions[r].height = h

# título
ws.merge_cells('C2:J3')
c = ws['C2']; c.value = 'Mis patrones saludables'; c.font = F(26, True, AZUL_TITULO); c.alignment = centro
ws.row_dimensions[2].height = 26; ws.row_dimensions[3].height = 22

# «Crea una pestaña llamada “Patrones v1”» y el icono de hoja verde
ws.merge_cells('D5:H5'); ws['D5'] = 'Crea una pestaña llamada'; ws['D5'].font = F(13); ws['D5'].alignment = izquierda
ws.merge_cells('D6:H6'); ws['D6'] = '“Patrones v1”'; ws['D6'].font = F(14, True); ws['D6'].alignment = izquierda
ws.merge_cells('I5:I6'); ws['I5'] = '📗'; ws['I5'].font = Font(size=26); ws['I5'].alignment = centro
ws.row_dimensions[5].height = 20; ws.row_dimensions[6].height = 22

# la tabla de Variable · Promedio · Máximo · Mínimo
ws.merge_cells('C8:E8'); ws['C8'] = 'Variable'
for ref, t in [('G8', 'Promedio'), ('H8', 'Máximo'), ('I8', 'Mínimo')]: ws[ref] = t
ws.merge_cells('G8:G8')
for ref in ['C8', 'G8', 'H8', 'I8']: ws[ref].font = F(12, True); ws[ref].alignment = centro
# E queda dentro de «Variable»: la columna del nombre se ensancha con D y E
ws.unmerge_cells('C8:E8'); ws.merge_cells('C8:F8')
filas = [('🥛', 'Agua', 'G'), ('😴', 'Sueño', 'H'), ('🏃', 'Actividad', 'I')]
for i, (ico, nombre, col) in enumerate(filas):
    r = 9 + i
    ws[f'C{r}'] = ico; ws[f'C{r}'].font = Font(size=18); ws[f'C{r}'].alignment = centro
    ws.merge_cells(f'D{r}:F{r}'); ws[f'D{r}'] = nombre; ws[f'D{r}'].font = F(12, True); ws[f'D{r}'].alignment = izquierda
    hay = f'COUNT({R(col)})'
    ws[f'G{r}'] = f'=IF({hay}=0,"",ROUND(AVERAGE({R(col)}),1))'
    ws[f'H{r}'] = f'=IF({hay}=0,"",MAX({R(col)}))'
    ws[f'I{r}'] = f'=IF({hay}=0,"",MIN({R(col)}))'
    for ref in [f'G{r}', f'H{r}', f'I{r}']: ws[ref].font = F(12); ws[ref].alignment = centro
    ws.row_dimensions[r].height = 30
ws.row_dimensions[8].height = 22
rejilla(ws, 'C8:I11')
marco(ws, 'C8:I11')

# comparativa (izquierda) y la gráfica (derecha)
ws.merge_cells('C13:E14')
ws['C13'] = 'Comparativa de promedios según nivel de energía'; ws['C13'].font = F(11, True); ws['C13'].alignment = izquierda
ws.merge_cells('F13:F14'); ws['F13'] = '🔋🧴'; ws['F13'].font = Font(size=14); ws['F13'].alignment = centro
ws.merge_cells('C15:E15'); ws['C15'] = 'Nivel de energía'; ws['F15'] = 'Agua'
for ref in ['C15', 'F15']: ws[ref].font = F(11, True); ws[ref].alignment = centro
for i, nivel in enumerate(['Bajo', 'Medio', 'Alto']):
    r = 16 + i
    ws.merge_cells(f'C{r}:E{r}'); ws[f'C{r}'] = nivel; ws[f'C{r}'].font = F(11); ws[f'C{r}'].alignment = izquierda
    n = f'COUNTIF({R("E")},"{nivel}")'
    ws[f'F{r}'] = f'=IF({n}=0,"",ROUND(AVERAGEIF({R("E")},"{nivel}",{R("G")}),1))'
    ws[f'F{r}'].font = F(11); ws[f'F{r}'].alignment = centro
    ws.row_dimensions[r].height = 26
rejilla(ws, 'C15:F18')
marco(ws, 'C13:F18')
for r in (13, 14, 15): ws.row_dimensions[r].height = 18 if r < 15 else 22

ws.merge_cells('G13:I14')
ws['G13'] = 'Dibuja aquí tu gráfica de barras 📊'; ws['G13'].font = F(10); ws['G13'].alignment = Alignment(horizontal='left', vertical='top', wrap_text=True, indent=1)
marco(ws, 'G13:I18')

graf = BarChart(); graf.type = 'col'; graf.style = 10
graf.title = None; graf.legend = None
graf.add_data(Reference(ws, min_col=6, min_row=15, max_row=18), titles_from_data=True)
graf.set_categories(Reference(ws, min_col=3, min_row=16, max_row=18))
graf.y_axis.title = 'Vasos'; graf.y_axis.majorGridlines = None
graf.dataLabels = DataLabelList(); graf.dataLabels.showVal = True
graf.dataLabels.showSerName = False; graf.dataLabels.showCatName = False; graf.dataLabels.showLegendKey = False
graf.series[0].graphicalProperties.solidFill = 'F4A623'
graf.height, graf.width = 3.5, 6.3
ws.add_chart(graf, 'G14')

# el patrón
ws.merge_cells('C20:I20')
ws['C20'] = 'El patrón que encontramos fue…'; ws['C20'].font = F(12, True); ws['C20'].alignment = izquierda
ws.row_dimensions[20].height = 24
# Con nivel de energía, el patrón sale de la comparativa. Sin él (el Forms no lo
# preguntó), sale del sueño: quienes duermen 7 h o más contra quienes duermen menos.
mucho = f'AVERAGEIF({R("H")},">=7",{R("G")})'
poco = f'AVERAGEIF({R("H")},"<7",{R("G")})'
por_sueno = (f'IF(OR(COUNTIF({R("H")},">=7")=0,COUNTIF({R("H")},"<7")=0),"",'
             f'IF(ROUND({mucho},1)>ROUND({poco},1),"Quienes duermen 7 horas o más toman más agua ("&ROUND({mucho},1)&" vasos) que quienes duermen menos ("&ROUND({poco},1)&").",'
             f'IF(ROUND({mucho},1)<ROUND({poco},1),"Quienes duermen menos de 7 horas toman más agua ("&ROUND({poco},1)&" vasos) que quienes duermen 7 o más ("&ROUND({mucho},1)&").",'
             f'"Duerman mucho o poco, toman la misma agua ("&ROUND({mucho},1)&" vasos).")))')
patron = ('=IF(COUNT(' + R('G') + ')<3,"",'
          'IF(OR(F16="",F18=""),' + por_sueno + ','
          'IF(F18>F16,"Quienes tienen energía ALTA toman más agua ("&F18&" vasos) que quienes la tienen BAJA ("&F16&").",'
          'IF(F18<F16,"Quienes tienen energía BAJA toman más agua ("&F16&" vasos) que quienes la tienen ALTA ("&F18&").",'
          '"Energía alta y baja toman la misma agua ("&F16&" vasos)."))))')
for i, r in enumerate((21, 22, 23)):
    ws.merge_cells(f'C{r}:I{r}')
    ws.row_dimensions[r].height = 36 if i == 0 else 24
    for col in 'CDEFGHI':
        ws[f'{col}{r}'].border = Border(bottom=Side(style='thin', color='9CA3AF'))
    ws[f'C{r}'].font = F(11, italica=(i == 0), color=GRIS if i == 0 else NEGRO); ws[f'C{r}'].alignment = izquierda
ws['C21'] = patron
marco(ws, 'C20:I24')
ws.row_dimensions[24].height = 8

# que se imprima en una hoja
ws.print_area = 'A1:L27'
ws.page_setup.orientation = 'portrait'
ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_options.horizontalCentered = True

wb.save(SALE)
print('listo:', SALE)
