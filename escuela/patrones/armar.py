"""armar.py · la hoja de «Mis patrones saludables» (Lección 2, Quantrix).

Arma un .xlsx con tres pestañas y se sube a Drive como Google Sheets:
  · «Encuesta»      — una fila por compañero (las columnas en el MISMO orden que
                      un Google Form vinculado: Marca temporal primero).
  · «Patrones v1»   — el formato que mandaron, con todo calculado solo:
                      promedio, máximo y mínimo; el agua según nivel de energía;
                      la gráfica de barras y el patrón escrito.
  · «Formulario»    — las preguntas tal cual para armar el Google Form.

⚠ Las fórmulas leen la encuesta con INDIRECT("'Encuesta'!…"): si se borra la
pestaña y se vincula un Form, basta con que la nueva se llame «Encuesta» y todo
vuelve a jalar. Con referencias normales, borrar la pestaña deja #REF! para
siempre aunque se cree otra con el mismo nombre.

No trae datos: los datos son de los compañeros. Sin respuestas, la hoja dice
lo que falta en vez de enseñar ceros.

  python3.13 escuela/patrones/armar.py  →  escuela/patrones/patrones-saludables.xlsx
"""
import os
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
SALE = os.path.join(AQUI, 'patrones-saludables.xlsx')
FILAS = 300                                   # cupo de respuestas

AZUL, AZUL_CLARO, VERDE, GRIS, TINTA = '1E6FD9', 'E8F1FC', '2E9E5B', 'F4F6F8', '1B2430'
fino = Side(style='thin', color='9AA7B4')
caja = Border(left=fino, right=fino, top=fino, bottom=fino)
cent = Alignment(horizontal='center', vertical='center', wrap_text=True)
izq = Alignment(horizontal='left', vertical='center', wrap_text=True)

def R(col):                                   # el rango de una columna de la encuesta
    return f'INDIRECT("\'Encuesta\'!{col}2:{col}{FILAS + 1}")'

def celda(ws, ref, valor=None, negrita=False, tam=11, color=TINTA, fondo=None, alinear=cent, borde=True):
    c = ws[ref]
    if valor is not None: c.value = valor
    c.font = Font(name='Arial', bold=negrita, size=tam, color=color)
    c.alignment = alinear
    if fondo: c.fill = PatternFill('solid', fgColor=fondo)
    if borde: c.border = caja
    return c

wb = Workbook()

# ── Encuesta ──────────────────────────────────────────────────────────────
en = wb.active; en.title = 'Encuesta'
cab = ['Marca temporal', 'Nombre', 'Agua (vasos al día)', 'Sueño (horas por noche)',
       'Actividad física (minutos al día)', 'Nivel de energía']
for i, t in enumerate(cab):
    c = en.cell(row=1, column=i + 1, value=t)
    c.font = Font(name='Arial', bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=AZUL)
    c.alignment = cent; c.border = caja
for col, w in zip('ABCDEF', [20, 26, 16, 18, 20, 18]): en.column_dimensions[col].width = w
en.row_dimensions[1].height = 34
en.freeze_panes = 'A2'
for r in range(2, 42):
    for col in 'ABCDEF':
        en[f'{col}{r}'].border = caja
        if r % 2: en[f'{col}{r}'].fill = PatternFill('solid', fgColor=GRIS)
v_agua = DataValidation(type='whole', operator='between', formula1='0', formula2='30', allow_blank=True,
                        error='Vasos de agua al día: un número entero de 0 a 30.', errorTitle='Agua', showErrorMessage=True)
v_sueno = DataValidation(type='decimal', operator='between', formula1='0', formula2='24', allow_blank=True,
                         error='Horas de sueño: de 0 a 24.', errorTitle='Sueño', showErrorMessage=True)
v_act = DataValidation(type='whole', operator='between', formula1='0', formula2='600', allow_blank=True,
                       error='Minutos de actividad al día: de 0 a 600.', errorTitle='Actividad', showErrorMessage=True)
v_ene = DataValidation(type='list', formula1='"Bajo,Medio,Alto"', allow_blank=True,
                       error='Elige Bajo, Medio o Alto.', errorTitle='Energía', showErrorMessage=True)
for v, col in [(v_agua, 'C'), (v_sueno, 'D'), (v_act, 'E'), (v_ene, 'F')]:
    en.add_data_validation(v); v.add(f'{col}2:{col}{FILAS + 1}')

# ── Patrones v1 ───────────────────────────────────────────────────────────
pa = wb.create_sheet('Patrones v1')
for col, w in zip('ABCDEFGH', [4, 26, 15, 15, 15, 4, 30, 4]): pa.column_dimensions[col].width = w

pa.merge_cells('B1:G1')
celda(pa, 'B1', 'Mis patrones saludables', negrita=True, tam=22, color=AZUL, borde=False)
pa.row_dimensions[1].height = 40
pa.merge_cells('B2:G2')
celda(pa, 'B2', '=IF(COUNTA(' + R('B') + ')=0,"Todavía no hay respuestas en la pestaña Encuesta.",'
      '"Respuestas registradas: "&COUNTA(' + R('B') + '))', color='5D6B78', borde=False)

for ref, t in zip(['B4', 'C4', 'D4', 'E4'], ['Variable', 'Promedio', 'Máximo', 'Mínimo']):
    celda(pa, ref, t, negrita=True, fondo=AZUL_CLARO)
filas = [('💧  Agua (vasos al día)', 'C'), ('😴  Sueño (horas)', 'D'), ('🏃  Actividad (minutos)', 'E')]
for i, (nombre, col) in enumerate(filas):
    r = 5 + i
    celda(pa, f'B{r}', nombre, negrita=True, alinear=izq)
    hay = f'COUNT({R(col)})'
    celda(pa, f'C{r}', f'=IF({hay}=0,"—",ROUND(AVERAGE({R(col)}),1))')
    celda(pa, f'D{r}', f'=IF({hay}=0,"—",MAX({R(col)}))')
    celda(pa, f'E{r}', f'=IF({hay}=0,"—",MIN({R(col)}))')
    pa.row_dimensions[r].height = 26

pa.merge_cells('B9:D9')
celda(pa, 'B9', 'Comparativa de promedios según nivel de energía 🔋', negrita=True, fondo=AZUL_CLARO, alinear=izq)
for ref, t in zip(['B10', 'C10', 'D10'], ['Nivel de energía', 'Agua (vasos)', 'Personas']):
    celda(pa, ref, t, negrita=True)
for i, nivel in enumerate(['Bajo', 'Medio', 'Alto']):
    r = 11 + i
    celda(pa, f'B{r}', nivel, alinear=izq)
    n = f'COUNTIF({R("F")},"{nivel}")'
    # vacío —no cero— cuando nadie eligió ese nivel: un 0 en la gráfica diría «no toman agua»
    celda(pa, f'C{r}', f'=IF({n}=0,"",ROUND(AVERAGEIF({R("F")},"{nivel}",{R("C")}),1))')
    celda(pa, f'D{r}', f'={n}')

graf = BarChart(); graf.type = 'col'; graf.style = 10
graf.title = 'Agua según nivel de energía'
graf.y_axis.title = 'Vasos al día (promedio)'; graf.x_axis.title = 'Nivel de energía'
graf.add_data(Reference(pa, min_col=3, min_row=10, max_row=13), titles_from_data=True)
graf.set_categories(Reference(pa, min_col=2, min_row=11, max_row=13))
graf.legend = None
graf.dataLabels = DataLabelList(); graf.dataLabels.showVal = True
# sólo el número: sin esto LibreOffice escribe «Bajo; Agua (vasos); 3.5» encima de cada barra
graf.dataLabels.showSerName = False; graf.dataLabels.showCatName = False; graf.dataLabels.showLegendKey = False
graf.series[0].graphicalProperties.solidFill = VERDE
graf.height, graf.width = 7.2, 10.5
pa.add_chart(graf, 'F9')

pa.merge_cells('B24:G24')
celda(pa, 'B24', 'El patrón que encontramos fue…', negrita=True, tam=13, fondo=AZUL_CLARO, alinear=izq)
patron = ('=IF(COUNTA(' + R('B') + ')<3,"Cuando haya al menos 3 respuestas, aquí sale el patrón.",'
          'IF(OR(C11="",C13=""),"Faltan respuestas de energía Alta o Baja para comparar.",'
          'IF(C13>C11,"Quienes se sienten con energía ALTA toman más agua ("&C13&" vasos al día) que quienes la sienten BAJA ("&C11&"): más agua va de la mano con más energía.",'
          'IF(C13<C11,"Quienes se sienten con energía BAJA toman más agua ("&C11&" vasos) que quienes la sienten ALTA ("&C13&"): en nuestro grupo el agua no explica la energía; hay que mirar el sueño y la actividad.",'
          '"Energía alta y baja toman la misma agua ("&C11&" vasos): aquí no se ve relación."))))')
pa.merge_cells('B25:G26')
celda(pa, 'B25', patron, alinear=izq)
pa.row_dimensions[25].height = 30; pa.row_dimensions[26].height = 30
pa.merge_cells('B27:G29')
celda(pa, 'B27', 'Con nuestras palabras: ', color='5D6B78', alinear=Alignment(horizontal='left', vertical='top', wrap_text=True))
pa.sheet_view.showGridLines = False

# ── Formulario ────────────────────────────────────────────────────────────
fo = wb.create_sheet('Formulario')
fo.column_dimensions['A'].width = 4; fo.column_dimensions['B'].width = 46; fo.column_dimensions['C'].width = 40
celda(fo, 'B1', 'Las preguntas para el Google Form', negrita=True, tam=15, color=AZUL, borde=False, alinear=izq)
celda(fo, 'B2', 'En este orden: así las respuestas caen en las mismas columnas que la pestaña Encuesta.',
      color='5D6B78', borde=False, alinear=izq)
celda(fo, 'B4', 'Pregunta', negrita=True, fondo=AZUL_CLARO); celda(fo, 'C4', 'Tipo en Forms', negrita=True, fondo=AZUL_CLARO)
preguntas = [
    ('¿Cómo te llamas?', 'Respuesta corta · obligatoria'),
    ('¿Cuántos vasos de agua tomas al día?', 'Respuesta corta · número entero de 0 a 30'),
    ('¿Cuántas horas duermes por noche?', 'Respuesta corta · número de 0 a 24'),
    ('¿Cuántos minutos de actividad física haces al día?', 'Respuesta corta · número de 0 a 600'),
    ('¿Cómo sientes tu nivel de energía?', 'Opción múltiple: Bajo · Medio · Alto'),
]
for i, (q, tipo) in enumerate(preguntas):
    celda(fo, f'B{5 + i}', f'{i + 1}. {q}', alinear=izq); celda(fo, f'C{5 + i}', tipo, alinear=izq)
pasos = ['Para juntar las respuestas del Form aquí:',
         '1. En el Form: Respuestas → Vincular con Hojas de cálculo → Seleccionar una hoja existente → esta.',
         '2. Aparece una pestaña «Respuestas de formulario 1».',
         '3. Borra la pestaña «Encuesta» vacía y cámbiale el nombre a la nueva por «Encuesta».',
         '4. Listo: «Patrones v1» se llena solo con cada respuesta.']
for i, t in enumerate(pasos):
    fo.merge_cells(f'B{12 + i}:C{12 + i}')
    celda(fo, f'B{12 + i}', t, negrita=(i == 0), borde=False, alinear=izq)
fo.sheet_view.showGridLines = False

wb.active = 1                                  # que abra en el formato
wb.save(SALE)
print('listo:', SALE)
